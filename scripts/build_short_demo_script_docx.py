from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK


OUT = r"D:\DATN\docs\LOI_THOAI_DEMO_FEA_6_NGUOI_30_PHUT_NGAN_GON.docx"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
PALE = "F5F8FC"
GRAY = "687386"
WHITE = "FFFFFF"
BLACK = "172033"


def set_font(run, name="Calibri", size=11, bold=False, color=BLACK, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths_dxa):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[i]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[i] / 1440)
            cell_margins(cell)


def no_split(paragraph):
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.widow_control = True


def add_timing(doc, time_text, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{time_text}  |  {title}")
    set_font(r, size=11, bold=True, color=DARK_BLUE)


def add_action(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.keep_with_next = True
    r1 = p.add_run("THAO TÁC: ")
    set_font(r1, size=9, bold=True, color=GRAY)
    r2 = p.add_run(text)
    set_font(r2, size=9, italic=True, color=GRAY)


def add_speech(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    no_split(p)
    r = p.add_run(f'“{text}”')
    set_font(r, size=11, color=BLACK)


def add_speech_parts(doc, parts):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    no_split(p)
    for index, (text, bold) in enumerate(parts):
        prefix = "“" if index == 0 else ""
        suffix = "”" if index == len(parts) - 1 else ""
        r = p.add_run(prefix + text + suffix)
        set_font(r, size=11, bold=bold, color=BLACK)


def add_note(doc, text, label="LƯU Ý"):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.1
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), PALE)
    p_pr.append(shd)
    r1 = p.add_run(f"{label}: ")
    set_font(r1, size=9.5, bold=True, color=DARK_BLUE)
    r2 = p.add_run(text)
    set_font(r2, size=9.5, color=BLACK)


def add_person_header(doc, number, time_text, role):
    if len(doc.paragraphs) > 0:
        doc.add_page_break()
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"NGƯỜI {number}")
    set_font(r, size=12, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 1"]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(role)
    set_font(r, size=19, bold=True, color=BLACK)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(9)
    r = p.add_run(time_text)
    set_font(r, size=10.5, bold=True, color=GRAY)


def add_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run("FEA Learning  •  Kịch bản luyện thuyết trình  •  30 phút")
    set_font(r, size=8.5, color=GRAY)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.3)
section.footer_distance = Inches(0.3)
add_footer(section)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25
for name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 14, 7),
    ("Heading 3", 12, DARK_BLUE, 10, 5),
):
    st = styles[name]
    st.font.name = "Calibri"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

# Cover / overview
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(20)
p.paragraph_format.space_after = Pt(3)
r = p.add_run("KỊCH BẢN DEMO FEA LEARNING")
set_font(r, size=24, bold=True, color=BLACK)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(18)
r = p.add_run("Bản lời thoại ngắn gọn để học • 6 người • 30 phút")
set_font(r, size=13, color=BLUE)

add_note(doc, "Chỉ đọc phần trong dấu ngoặc kép. Dòng THAO TÁC chỉ để người trình chiếu làm theo.", "CÁCH DÙNG")

table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.LEFT
table.style = "Table Grid"
headers = ["Người", "Thời gian", "Nội dung"]
for i, text in enumerate(headers):
    cell = table.rows[0].cells[i]
    shade(cell, LIGHT_BLUE)
    p = cell.paragraphs[0]
    r = p.add_run(text)
    set_font(r, size=10, bold=True, color=DARK_BLUE)
tr_pr = table.rows[0]._tr.get_or_add_trPr()
tbl_header = OxmlElement("w:tblHeader")
tbl_header.set(qn("w:val"), "true")
tr_pr.append(tbl_header)
rows = [
    ("1", "00:00–07:00", "Slide và trang chủ"),
    ("2", "07:00–12:00", "Tìm, mua và thanh toán"),
    ("3", "12:00–16:30", "Học, học liệu, quiz, tiến độ"),
    ("4", "16:30–21:00", "Đăng ký và Admin duyệt giảng viên"),
    ("5", "21:00–25:00", "Tạo và Admin duyệt khóa học"),
    ("6", "25:00–30:00", "Rút tiền và Admin duyệt chi trả"),
]
for values in rows:
    cells = table.add_row().cells
    for i, text in enumerate(values):
        p = cells[i].paragraphs[0]
        r = p.add_run(text)
        set_font(r, size=10)
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
set_table_widths(table, [900, 1900, 6560])

doc.add_paragraph()
p = doc.add_paragraph()
p.style = styles["Heading 2"]
r = p.add_run("Nguyên tắc nói đúng")
set_font(r, size=13, bold=True, color=BLUE)
for text in [
    "Chỉ nói thành công sau khi màn hình hiện kết quả.",
    "Tạo QR chưa có nghĩa là đã thanh toán.",
    "Duyệt giảng viên và duyệt khóa học là hai quy trình riêng.",
    "Duyệt rút tiền trong ứng dụng không tự chứng minh ngân hàng đã chuyển khoản.",
]:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_font(r, size=10.5)

# Person 1
add_person_header(doc, 1, "00:00–07:00  •  7 phút", "Giới thiệu và trang chủ")
add_timing(doc, "00:00–01:30", "Mở đầu và bài toán")
add_speech(doc, "Kính thưa quý thầy cô, nhóm em xin trình bày FEA Learning, hệ thống học trực tuyến dành cho học viên, giảng viên và quản trị viên.")
add_speech(doc, "Hệ thống tập trung khóa học, nội dung, tiến độ, thanh toán và kiểm duyệt trên cùng một nền tảng.")
add_timing(doc, "01:30–03:15", "Mục tiêu và chức năng")
add_speech(doc, "Học viên có thể tìm, mua, học và theo dõi tiến độ. Giảng viên quản lý chương, bài học, học liệu và quiz. Admin quản lý tài khoản, nội dung và giao dịch.")
add_timing(doc, "03:15–05:00", "Quy trình và công nghệ")
add_speech(doc, "Luồng demo hôm nay đi từ mua và học khóa học, đăng ký giảng viên, tạo khóa học, phê duyệt, đến doanh thu và rút tiền.")
add_speech(doc, "Dự án được xây dựng bằng Laravel, dữ liệu được xử lý phía máy chủ và phân quyền theo từng vai trò.")
add_timing(doc, "05:00–06:45", "Demo trang chủ")
add_action(doc, "Mở trang chủ; chỉ logo, tìm kiếm, danh mục và thẻ khóa học.")
add_speech(doc, "Đây là trang chủ FEA Learning. Người dùng có thể tìm kiếm theo khóa học, kỹ năng hoặc giảng viên, sau đó xem danh mục và các khóa đang được công khai.")
add_speech(doc, "Khách có thể xem thông tin giới thiệu; các thao tác mua và học yêu cầu đăng nhập.")
add_timing(doc, "06:45–07:00", "Bàn giao")
add_speech(doc, "Tiếp theo, bạn [người 2] sẽ vào vai học viên và thực hiện mua khóa học.")

# Person 2
add_person_header(doc, 2, "07:00–12:00  •  5 phút", "Học viên tìm và mua khóa học")
add_timing(doc, "07:00–07:35", "Đăng nhập")
add_action(doc, "Đăng nhập bằng tài khoản học viên đã chuẩn bị.")
add_speech(doc, "Em đăng nhập bằng tài khoản học viên. Sau khi xác thực, hệ thống hiển thị các chức năng đúng với vai trò này.")
add_timing(doc, "07:35–08:25", "Tìm và lọc")
add_action(doc, "Tìm [từ khóa], chọn bộ lọc, mở khóa [tên khóa].")
add_speech(doc, "Em tìm khóa học theo từ khóa và dùng bộ lọc để thu hẹp kết quả.")
add_timing(doc, "08:25–09:15", "Xem chi tiết")
add_speech(doc, "Trang chi tiết cho biết tên khóa, giảng viên, giá, mục tiêu và chương trình học trước khi mua.")
add_timing(doc, "09:15–10:10", "Tạo đơn")
add_action(doc, "Thêm vào giỏ hoặc Mua ngay; kiểm tra tổng tiền.")
add_speech(doc, "Em chọn khóa học, kiểm tra lại nội dung và tổng tiền, sau đó chuyển sang thanh toán.")
add_timing(doc, "10:10–11:20", "Thanh toán")
add_action(doc, "Chọn PayOS/VietQR hoặc MoMo; tạo QR và chờ xác nhận.")
add_speech(doc, "Hệ thống đã tạo yêu cầu thanh toán cho đơn [mã đơn]. Tạo QR chưa có nghĩa giao dịch đã thành công.")
add_speech(doc, "Sau khi nhận xác nhận, hệ thống cập nhật đơn và cấp quyền học.")
add_timing(doc, "11:20–12:00", "Kiểm tra và bàn giao")
add_action(doc, "Chỉ nói câu sau nếu màn hình báo thành công; bấm Vào học ngay.")
add_speech(doc, "Giao dịch đã được xác nhận. Khóa học hiện có thể mở bằng tài khoản này. Bạn [người 3] sẽ tiếp tục phần học.")
add_note(doc, "Nếu thanh toán chưa xác nhận: “Nhóm em chưa kết luận thành công và chuyển sang tài khoản đã mua sẵn để tiếp tục.”", "CÂU DỰ PHÒNG")

# Person 3
add_person_header(doc, 3, "12:00–16:30  •  4 phút 30 giây", "Học bài, học liệu, quiz và tiến độ")
add_timing(doc, "12:00–12:40", "Tổng quan khóa học")
add_action(doc, "Mở trang học; chỉ nội dung và danh sách chương.")
add_speech(doc, "Đây là giao diện học gồm nội dung hiện tại và danh sách chương, bài của khóa học.")
add_timing(doc, "12:40–13:30", "Học một bài")
add_action(doc, "Phát video 10–15 giây hoặc mở bài đọc.")
add_speech(doc, "Em mở một bài học. Tiến độ được ghi nhận trong quá trình học, không chỉ dựa vào việc mở trang hoặc kéo video đến cuối.")
add_timing(doc, "13:30–14:10", "Học liệu và bài tập")
add_action(doc, "Chỉ tab học liệu và một bài tập đã chuẩn bị.")
add_speech(doc, "Bài học có thể kèm tài liệu và bài tập để học viên tải về, xem yêu cầu và nộp kết quả.")
add_timing(doc, "14:10–15:30", "Làm quiz")
add_action(doc, "Mở quiz ngắn, chọn đáp án và nộp.")
add_speech(doc, "Em làm quiz và nộp bài. Hệ thống chấm theo cấu hình của bài kiểm tra.")
add_speech(doc, "Kết quả lượt này là [điểm thực tế], trạng thái [đạt/chưa đạt].")
add_timing(doc, "15:30–16:30", "Tiến độ và bàn giao")
add_action(doc, "Mở tiến độ khóa học.")
add_speech(doc, "Học viên có thể xem phần đã hoàn thành và phần còn lại của khóa học.")
add_speech(doc, "Tiếp theo, bạn [người 4] sẽ trình bày quy trình đăng ký và phê duyệt giảng viên.")

# Person 4
add_person_header(doc, 4, "16:30–21:00  •  4 phút 30 giây", "Đăng ký và Admin duyệt giảng viên")
add_timing(doc, "16:30–17:30", "Tạo tài khoản")
add_action(doc, "Tab 1: mở form đăng ký giảng viên; điền dữ liệu đã chuẩn bị.")
add_speech(doc, "Em đăng ký tài khoản giảng viên và nhập các thông tin theo yêu cầu của hệ thống.")
add_timing(doc, "17:30–18:30", "Hoàn thiện hồ sơ")
add_action(doc, "Chọn lĩnh vực, bổ sung giới thiệu và tài liệu minh chứng.")
add_speech(doc, "Em bổ sung phần giới thiệu, lĩnh vực giảng dạy và tài liệu minh chứng.")
add_speech(doc, "Lĩnh vực này còn được dùng để kiểm tra danh mục khóa học mà giảng viên được phép quản lý.")
add_timing(doc, "18:30–19:00", "Gửi hồ sơ")
add_action(doc, "Gửi hồ sơ; chỉ trạng thái thực tế.")
add_speech(doc, "Em gửi hồ sơ cho admin. Trạng thái hiện tại là [trạng thái thực tế]; gửi thành công chưa có nghĩa đã được duyệt.")
add_timing(doc, "19:00–20:20", "Admin kiểm tra và duyệt")
add_action(doc, "Tab 2/profile Admin: mở Đơn đăng ký giảng viên, tìm đúng email, kiểm tra và bấm Duyệt.")
add_speech(doc, "Ở phía admin, em mở hồ sơ vừa gửi và kiểm tra thông tin, lĩnh vực cùng tài liệu minh chứng.")
add_speech(doc, "Hồ sơ đáp ứng điều kiện demo nên em thực hiện phê duyệt.")
add_timing(doc, "20:20–21:00", "Xác nhận kết quả")
add_action(doc, "Chỉ nói khi trạng thái đã cập nhật; quay lại tab giảng viên và tải lại.")
add_speech(doc, "Hệ thống đã cập nhật hồ sơ sang trạng thái [trạng thái thực tế]. Đây là duyệt giảng viên, chưa phải duyệt khóa học.")
add_speech(doc, "Bạn [người 5] sẽ tiếp tục tạo và gửi một khóa học để admin kiểm duyệt riêng.")

# Person 5
add_person_header(doc, 5, "21:00–25:00  •  4 phút", "Tạo và Admin duyệt khóa học")
add_timing(doc, "21:00–21:40", "Tạo khóa học")
add_action(doc, "Tab giảng viên: tạo khóa [tên khóa], chọn đúng danh mục.")
add_speech(doc, "Em tạo khóa học mới, nhập thông tin cơ bản và chọn danh mục phù hợp với lĩnh vực đã được duyệt.")
add_timing(doc, "21:40–22:30", "Tạo chương và bài")
add_action(doc, "Thêm một chương và bài đọc ngắn; chỉ học liệu/quiz đã chuẩn bị.")
add_speech(doc, "Khóa học được tổ chức theo chương và bài. Hệ thống hỗ trợ video, bài đọc, tài liệu, bài tập và quiz.")
add_timing(doc, "22:30–23:05", "Gửi duyệt")
add_action(doc, "Dùng khóa đã hoàn thiện checklist; bấm Gửi duyệt.")
add_speech(doc, "Em kiểm tra checklist và gửi khóa học cho admin. Khóa hiện ở trạng thái [trạng thái thực tế] và chưa tự động được công khai.")
add_timing(doc, "23:05–24:20", "Admin kiểm duyệt")
add_action(doc, "Tab/profile Admin: mở Khóa học chờ duyệt, kiểm tra rồi bấm Duyệt khóa học.")
add_speech(doc, "Admin kiểm tra thông tin, chương trình, bài học và điều kiện của khóa học.")
add_speech(doc, "Khóa đáp ứng điều kiện demo nên em thực hiện phê duyệt.")
add_timing(doc, "24:20–25:00", "Kết quả và bàn giao")
add_speech(doc, "Hệ thống đã cập nhật khóa sang trạng thái [trạng thái thực tế]. Duyệt giảng viên và duyệt khóa học là hai quy trình độc lập.")
add_speech(doc, "Tiếp theo, bạn [người 6] sẽ trình bày doanh thu, rút tiền và bước admin xử lý yêu cầu.")

# Person 6
add_person_header(doc, 6, "25:00–30:00  •  5 phút", "Giảng viên rút tiền và Admin duyệt chi trả")
add_timing(doc, "25:00–25:45", "Xem doanh thu")
add_action(doc, "Tab giảng viên: mở Doanh thu; chỉ tổng thu nhập, đã rút và số dư khả dụng.")
add_speech(doc, "Sau khi có đơn hàng đã thanh toán, hệ thống ghi nhận doanh thu cho giảng viên.")
add_speech_parts(doc, [
    ("Sau khi học viên thanh toán khóa học thành công, hệ thống sẽ ghi nhận doanh thu cho giảng viên. Trong đó, giảng viên nhận được ", False),
    ("80% giá trị khóa học", True),
    (", còn ", False),
    ("20% được khấu trừ làm phí dịch vụ của nền tảng", True),
    (". Giảng viên có thể theo dõi doanh thu và số tiền thực nhận tại khu vực quản lý tài chính.", False),
])
add_speech(doc, "Giảng viên có thể xem tổng thu nhập, khoản đã rút và số dư khả dụng.")
add_timing(doc, "25:45–26:40", "Gửi yêu cầu rút tiền")
add_action(doc, "Nhập số tiền và thông tin ngân hàng đã chuẩn bị; gửi yêu cầu.")
add_speech(doc, "Em nhập số tiền [số tiền] và kiểm tra thông tin nhận tiền.")
add_speech(doc, "Yêu cầu đã được tạo và đang chờ admin xử lý; ở bước này chưa thể coi là đã chi trả.")
add_timing(doc, "26:40–27:50", "Admin kiểm tra")
add_action(doc, "Chuyển profile Admin; mở Yêu cầu rút tiền và tìm đúng yêu cầu.")
add_speech(doc, "Ở phía admin, em kiểm tra tên giảng viên, số tiền, thông tin ngân hàng và nguồn tiền khả dụng.")
add_speech(doc, "Các thông tin này khớp với yêu cầu vừa được gửi.")
add_timing(doc, "27:50–28:40", "Admin duyệt")
add_action(doc, "Bấm duyệt và chờ thông báo; chỉ nói kết quả sau khi trạng thái cập nhật.")
add_speech(doc, "Yêu cầu đáp ứng điều kiện nên em thực hiện phê duyệt.")
add_speech(doc, "Hệ thống đã cập nhật yêu cầu sang trạng thái [trạng thái thực tế].")
add_timing(doc, "28:40–29:20", "Kiểm tra phía giảng viên")
add_action(doc, "Quay lại tab giảng viên, tải lại lịch sử rút tiền.")
add_speech(doc, "Phía giảng viên hiện đã thấy kết quả xử lý của yêu cầu vừa gửi.")
add_timing(doc, "29:20–30:00", "Tổng kết")
add_speech(doc, "Qua phần demo, nhóm em đã trình bày luồng học viên, giảng viên và admin trên cùng một hệ thống.")
add_speech(doc, "Phần trình bày của nhóm em xin kết thúc. Nhóm em cảm ơn quý thầy cô đã lắng nghe.")
add_note(doc, "Nên nói “Admin duyệt yêu cầu rút tiền”. Không nói “hệ thống đã tự động chuyển tiền vào ngân hàng” nếu chưa có bằng chứng giao dịch bên ngoài.", "CÂU NÓI AN TOÀN")

# Final quick checklist page
doc.add_page_break()
p = doc.add_paragraph()
p.style = styles["Heading 1"]
p.paragraph_format.space_before = Pt(0)
r = p.add_run("Checklist luyện tập trước buổi bảo vệ")
set_font(r, size=18, bold=True, color=BLACK)
items = [
    "Mỗi vai trò dùng một profile trình duyệt riêng.",
    "Học viên chưa sở hữu khóa dùng để mua; có tài khoản dự phòng đã mua sẵn.",
    "Quiz ngắn, khóa học và video đã được kiểm tra trước.",
    "Người 4 ghi lại email hồ sơ để Admin tìm nhanh.",
    "Người 5 dùng khóa đã đủ checklist trước khi gửi duyệt.",
    "Người 6 dùng giảng viên có đủ số dư khả dụng và chưa có yêu cầu trùng.",
    "Chỉ đọc kết quả thực tế đang hiện trên màn hình.",
    "Có người bấm giờ và báo khi còn 30 giây.",
]
for text in items:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text)
    set_font(r, size=11)
add_note(doc, "Nếu có lỗi hoặc chờ quá lâu, công bố dùng dữ liệu dự phòng. Không sửa quyền, bỏ kiểm tra hoặc nhận một kết quả chưa xảy ra.", "NGUYÊN TẮC")

# Keep headers and footers consistent.
for sec in doc.sections:
    add_footer(sec)

doc.core_properties.title = "Lời thoại demo FEA Learning - 6 người - 30 phút"
doc.core_properties.subject = "Kịch bản thuyết trình ngắn gọn để luyện tập"
doc.core_properties.author = "Nhóm FEA Learning"
doc.save(OUT)
print(OUT)
