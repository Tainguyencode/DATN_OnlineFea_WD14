from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"D:\DATN\docs\LOI_THOAI_CHI_TIET_NGUOI_6_RUT_TIEN_VA_ADMIN_DUYET.docx"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
GRAY = "687386"
BLACK = "172033"
PALE = "F5F8FC"


def font(run, size=11, bold=False, italic=False, color=BLACK):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    font(r, size=16 if level == 1 else 13, bold=True, color=BLUE if level == 1 else DARK_BLUE)
    return p


def timing(doc, time_text, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{time_text}  |  {title}")
    font(r, size=11.5, bold=True, color=DARK_BLUE)


def action(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run("THAO TÁC: ")
    font(r, size=9, bold=True, color=GRAY)
    r = p.add_run(text)
    font(r, size=9, italic=True, color=GRAY)


def speech(doc, text, bold_phrases=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_together = True
    if not bold_phrases:
        r = p.add_run(f'“{text}”')
        font(r)
        return
    remaining = text
    first = True
    for phrase in bold_phrases:
        before, sep, after = remaining.partition(phrase)
        if sep:
            r = p.add_run(("“" if first else "") + before)
            font(r)
            first = False
            r = p.add_run(sep)
            font(r, bold=True)
            remaining = after
    r = p.add_run(("“" if first else "") + remaining + "”")
    font(r)


def note(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.1
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), PALE)
    p._p.get_or_add_pPr().append(shd)
    r = p.add_run(label + ": ")
    font(r, size=9.5, bold=True, color=DARK_BLUE)
    r = p.add_run(text)
    font(r, size=9.5)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(1)
sec.right_margin = Inches(1)
sec.header_distance = Inches(0.3)
sec.footer_distance = Inches(0.3)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25
for name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, DARK_BLUE, 14, 7),
):
    st = doc.styles[name]
    st.font.name = "Calibri"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("FEA Learning  •  Người 6  •  Doanh thu và rút tiền")
font(r, size=8.5, color=GRAY)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(16)
p.paragraph_format.space_after = Pt(3)
r = p.add_run("LỜI THOẠI CHI TIẾT - NGƯỜI THỨ 6")
font(r, size=23, bold=True)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(12)
r = p.add_run("Giảng viên rút tiền và Admin duyệt yêu cầu • 25:00–30:00")
font(r, size=13, color=BLUE)
note(doc, "CHUẨN BỊ", "Mở hai profile trình duyệt riêng: một tài khoản giảng viên có đủ số dư khả dụng và một tài khoản Admin.")

timing(doc, "25:00–25:20", "Giới thiệu")
action(doc, "Đứng tại trang quản lý tài chính của giảng viên.")
speech(doc, "Em xin tiếp tục phần cuối của hệ thống với quy trình quản lý doanh thu và rút tiền dành cho giảng viên.")
speech(doc, "Phần demo gồm hai bước: giảng viên gửi yêu cầu rút tiền và quản trị viên kiểm tra, phê duyệt yêu cầu đó.")

timing(doc, "25:20–26:00", "Giải thích cách chia doanh thu")
action(doc, "Chỉ phần tổng doanh thu hoặc chi tiết giao dịch.")
speech(doc, "Sau khi học viên thanh toán khóa học thành công, hệ thống sẽ ghi nhận doanh thu cho giảng viên.")
speech(doc, "Theo tỷ lệ mặc định đang áp dụng, giảng viên nhận được 80% giá trị khóa học, còn 20% được khấu trừ làm phí dịch vụ của nền tảng.", ["80% giá trị khóa học", "20% được khấu trừ làm phí dịch vụ của nền tảng"])
speech(doc, "Giảng viên có thể theo dõi doanh thu và số tiền thực nhận tại khu vực quản lý tài chính.")
speech(doc, "Ví dụ, với khóa học có giá một triệu đồng, giảng viên nhận tám trăm nghìn đồng và nền tảng nhận hai trăm nghìn đồng.")

timing(doc, "26:00–26:30", "Giới thiệu số liệu tài chính")
action(doc, "Chỉ tổng thu nhập, đã rút, đang chờ và số dư khả dụng.")
speech(doc, "Tại đây, giảng viên có thể theo dõi tổng thu nhập, số tiền đã rút, số tiền đang chờ xử lý và số dư khả dụng.")
speech(doc, "Tổng thu nhập là doanh thu đã được ghi nhận. Số dư khả dụng là số tiền giảng viên hiện có thể yêu cầu rút.")
speech(doc, "Nếu đã có yêu cầu đang chờ xử lý, khoản tiền đó sẽ không tiếp tục được sử dụng để tạo một yêu cầu trùng.")
note(doc, "NẾU DÙNG DỮ LIỆU CHUẨN BỊ", "Nói: “Các số liệu trên màn hình bao gồm dữ liệu demo đã chuẩn bị trước, không chỉ riêng giao dịch vừa thực hiện ở phần học viên.”")

timing(doc, "26:30–27:05", "Tạo yêu cầu rút tiền")
action(doc, "Mở form, nhập số tiền và thông tin ngân hàng đã chuẩn bị.")
speech(doc, "Bây giờ, em thực hiện tạo một yêu cầu rút tiền.")
speech(doc, "Em nhập số tiền muốn rút là [số tiền thực tế] và kiểm tra lại thông tin nhận tiền.")
speech(doc, "Khi gửi yêu cầu, hệ thống sẽ kiểm tra số tiền có hợp lệ và có vượt quá số dư khả dụng hay không.")
action(doc, "Bấm gửi; chờ thông báo thành công rồi chỉ dòng yêu cầu mới.")
speech(doc, "Hệ thống đã tiếp nhận yêu cầu rút tiền. Trạng thái hiện tại là ‘Đang chờ xử lý’.")
speech(doc, "Ở bước này, yêu cầu mới chỉ được gửi cho Admin; tiền chưa được xem là đã chi trả.")

doc.add_page_break()
heading(doc, "Admin kiểm tra và phê duyệt", 1)

timing(doc, "27:05–27:20", "Chuyển sang Admin")
action(doc, "Chuyển sang profile Admin đã đăng nhập sẵn.")
speech(doc, "Tiếp theo, em chuyển sang vai trò quản trị viên để xử lý chính yêu cầu vừa được gửi.")
speech(doc, "Hai tài khoản được mở bằng hai profile trình duyệt riêng để không làm thay đổi phiên đăng nhập của nhau.")

timing(doc, "27:20–27:55", "Tìm yêu cầu")
action(doc, "Mở danh sách yêu cầu rút tiền và tìm đúng giảng viên.")
speech(doc, "Tại khu vực quản lý yêu cầu rút tiền, Admin có thể xem người gửi, số tiền và trạng thái của từng yêu cầu.")
speech(doc, "Đây là yêu cầu của giảng viên [tên giảng viên], với số tiền [số tiền].")
speech(doc, "Thông tin này khớp với yêu cầu vừa được tạo và hiện đang ở trạng thái chờ duyệt.")

timing(doc, "27:55–28:25", "Kiểm tra điều kiện")
action(doc, "Mở chi tiết yêu cầu; chỉ thông tin nhận tiền và nguồn tiền.")
speech(doc, "Trước khi duyệt, Admin kiểm tra tên giảng viên, số tiền yêu cầu và thông tin nhận tiền.")
speech(doc, "Hệ thống cũng kiểm tra lại nguồn tiền khả dụng tại thời điểm xử lý.")
speech(doc, "Việc kiểm tra lại giúp hạn chế phê duyệt một yêu cầu không còn đủ điều kiện.")
speech(doc, "Trong trường hợp demo này, thông tin hợp lệ và nguồn tiền vẫn đáp ứng yêu cầu.")

timing(doc, "28:25–28:55", "Phê duyệt")
action(doc, "Bấm Duyệt và chờ hệ thống phản hồi.")
speech(doc, "Sau khi kiểm tra đầy đủ thông tin, em thực hiện phê duyệt yêu cầu rút tiền.")
speech(doc, "Hệ thống đã xử lý thành công và cập nhật yêu cầu sang trạng thái [trạng thái thực tế].")
speech(doc, "Kết quả này được lưu lại để cả Admin và giảng viên cùng theo dõi.")
note(doc, "NẾU HIỂN THỊ ĐÃ DUYỆT", "Nói: “Yêu cầu hiện đã chuyển từ trạng thái ‘Đang chờ xử lý’ sang ‘Đã duyệt’.”")

timing(doc, "28:55–29:20", "Kiểm tra phía giảng viên")
action(doc, "Quay lại profile giảng viên và tải lại lịch sử rút tiền.")
speech(doc, "Em quay lại tài khoản giảng viên và tải lại dữ liệu để kiểm tra kết quả.")
speech(doc, "Yêu cầu vừa gửi hiện đã được cập nhật sang trạng thái [trạng thái thực tế].")
speech(doc, "Lịch sử rút tiền đã ghi nhận số tiền, thời gian và kết quả xử lý của yêu cầu.")

timing(doc, "29:20–29:40", "Tổng kết quy trình")
speech(doc, "Giảng viên theo dõi doanh thu và gửi yêu cầu trong giới hạn số dư khả dụng. Admin tiếp nhận, kiểm tra và phê duyệt yêu cầu.")
speech(doc, "Việc tách hai vai trò giúp giảng viên không thể tự phê duyệt yêu cầu rút tiền của chính mình.")

timing(doc, "29:40–30:00", "Kết thúc bài demo")
speech(doc, "Qua toàn bộ phần demo, nhóm em đã trình bày các quy trình chính của học viên, giảng viên và Admin trên FEA Learning.")
speech(doc, "Phần trình bày của nhóm em xin kết thúc tại đây. Nhóm em xin cảm ơn quý thầy cô đã lắng nghe và mong nhận được các câu hỏi, góp ý.")

doc.add_page_break()
heading(doc, "Câu trả lời khi hội đồng hỏi", 1)
heading(doc, "Vì sao chia 80%–20%?", 2)
speech(doc, "20% là tỷ lệ phí nền tảng mặc định, vì vậy giảng viên mặc định nhận 80%. Hệ thống cũng hỗ trợ Admin cấu hình tỷ lệ riêng cho từng giảng viên khi cần.")
heading(doc, "Tiền đã được chuyển thật chưa?", 2)
speech(doc, "Hệ thống hiện quản lý doanh thu, số dư, yêu cầu rút tiền và trạng thái phê duyệt. Việc chuyển khoản thực tế vẫn cần được thực hiện và đối soát với ngân hàng. Trạng thái ‘Đã duyệt’ không tự nó chứng minh giao dịch ngân hàng đã hoàn tất.")
heading(doc, "Nếu không thể duyệt yêu cầu", 2)
speech(doc, "Tại thời điểm kiểm tra, yêu cầu chưa đáp ứng điều kiện về nguồn tiền nên hệ thống không cho phép Admin phê duyệt. Nhóm em không bỏ qua điều kiện chỉ để tạo thông báo thành công.")
heading(doc, "Câu phải tránh", 2)
note(doc, "KHÔNG NÓI", "“Hệ thống đã tự động chuyển tiền vào tài khoản ngân hàng của giảng viên.”")
note(doc, "NÊN NÓI", "“Admin đã duyệt yêu cầu rút tiền và hệ thống đã ghi nhận trạng thái xử lý.”")

heading(doc, "Checklist trước khi trình bày", 1)
for text in [
    "Tài khoản giảng viên có đủ số dư khả dụng.",
    "Chưa có yêu cầu rút tiền đang chờ bị trùng.",
    "Thông tin ngân hàng đã được chuẩn bị và kiểm tra.",
    "Profile giảng viên và Admin được mở riêng.",
    "Chỉ nói thành công sau khi màn hình hiện kết quả.",
    "Đọc đúng số tiền và trạng thái thực tế trên màn hình.",
]:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    font(r)

doc.core_properties.title = "Lời thoại chi tiết người thứ 6 - FEA Learning"
doc.core_properties.subject = "Giảng viên rút tiền và Admin duyệt yêu cầu"
doc.core_properties.author = "Nhóm FEA Learning"
doc.save(OUT)
print(OUT)
